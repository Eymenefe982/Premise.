import customtkinter as ctk
import threading
import sqlite3
from tkinter import filedialog
from docx import Document
from rag_engine import run_literature_review

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

class LiteratureApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("AI Literature Review Assistant - Pro Edition")
        self.geometry("950x850")
        
        self.current_report = ""
        self.current_articles = []
        self.init_db()
        
        # Üst Başlık
        self.title_label = ctk.CTkLabel(self, text="PubMed & Gemini RAG Laboratory", font=ctk.CTkFont(size=24, weight="bold"))
        self.title_label.pack(pady=(15, 10))
        
        # Seçenekler Paneli
        self.options_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.options_frame.pack(pady=5, fill="x", padx=20)
        
        # 1. Satır: Yazar ve Dil
        ctk.CTkLabel(self.options_frame, text="Author:").grid(row=0, column=0, padx=5, pady=5)
        self.author_entry = ctk.CTkEntry(self.options_frame, placeholder_text="e.g. Smith J", width=140)
        self.author_entry.grid(row=0, column=1, padx=5, pady=5)
        
        ctk.CTkLabel(self.options_frame, text="Language:").grid(row=0, column=2, padx=5, pady=5)
        self.lang_dropdown = ctk.CTkOptionMenu(self.options_frame, values=["English", "Türkçe", "Español", "Deutsch"], width=120)
        self.lang_dropdown.grid(row=0, column=3, padx=5, pady=5)
        
        # 2. Satır: Makale Slider'ı
        ctk.CTkLabel(self.options_frame, text="Max Articles:").grid(row=1, column=0, padx=5, pady=10)
        self.article_slider = ctk.CTkSlider(self.options_frame, from_=5, to=30, number_of_steps=25, command=self.update_slider_lbl)
        self.article_slider.grid(row=1, column=1, columnspan=2, sticky="ew", padx=5)
        self.article_slider.set(15)
        
        self.slider_val_label = ctk.CTkLabel(self.options_frame, text="15", width=30)
        self.slider_val_label.grid(row=1, column=3, sticky="w")
        
        # Arama Çubuğu
        self.search_entry = ctk.CTkEntry(self, placeholder_text="Enter your research topic...", width=700, height=45)
        self.search_entry.pack(pady=10)
        
        # İlerleme Çubuğu ve Durum
        self.progress_bar = ctk.CTkProgressBar(self, width=700)
        self.progress_bar.pack(pady=5)
        self.progress_bar.set(0)
        
        self.status_label = ctk.CTkLabel(self, text="Ready.", text_color="gray")
        self.status_label.pack()
        
        # Butonlar için Alt Frame (Search ve History yan yana)
        self.action_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.action_frame.pack(pady=5)
        
        self.search_button = ctk.CTkButton(self.action_frame, text="Search Literature", command=self.start_search_thread, width=200, height=40)
        self.search_button.grid(row=0, column=0, padx=10)
        
        # YENİ: Geçmiş Butonu
        self.history_button = ctk.CTkButton(self.action_frame, text="🕒 History", command=self.open_history_window, 
                                            width=120, height=40, fg_color="#4a4a4a", hover_color="#5e5e5e")
        self.history_button.grid(row=0, column=1, padx=10)
        
        # Sonuç Kutusu
        self.result_textbox = ctk.CTkTextbox(self, width=850, height=400, wrap="word", font=ctk.CTkFont(size=14))
        self.result_textbox.pack(pady=10, padx=20, expand=True, fill="both")
        
        # Dışa Aktarma Butonları
        self.export_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.export_frame.pack(pady=(0, 20))
        
        self.word_btn = ctk.CTkButton(self.export_frame, text="Export to Word", fg_color="#2b5797", state="disabled", command=self.export_word)
        self.word_btn.grid(row=0, column=0, padx=10)
        
        self.ris_btn = ctk.CTkButton(self.export_frame, text="Export Zotero (.ris)", fg_color="#b91d47", state="disabled", command=self.export_ris)
        self.ris_btn.grid(row=0, column=1, padx=10)

    def init_db(self):
        """Yerel arama geçmişi veritabanını oluşturur."""
        self.conn = sqlite3.connect("history.db", check_same_thread=False)
        self.conn.execute("CREATE TABLE IF NOT EXISTS searches (id INTEGER PRIMARY KEY, date TEXT, query TEXT, author TEXT)")
        self.conn.commit()

    # --- YENİ EKLENEN GEÇMİŞ FONKSİYONLARI ---
    def open_history_window(self):
        """Geçmiş aramaları gösteren pencereyi açar."""
        history_win = ctk.CTkToplevel(self)
        history_win.title("Search History")
        history_win.geometry("650x450")
        history_win.attributes("-topmost", True) # Pencerenin hep üstte kalmasını sağlar
        
        ctk.CTkLabel(history_win, text="Recent Searches", font=ctk.CTkFont(size=18, weight="bold")).pack(pady=15)
        
        # Scroll edilebilir liste alanı
        scroll_frame = ctk.CTkScrollableFrame(history_win, width=600, height=350)
        scroll_frame.pack(pady=10, padx=10, fill="both", expand=True)
        
        cursor = self.conn.cursor()
        # En son yapılan 50 aramayı getir
        cursor.execute("SELECT date, query, author FROM searches ORDER BY id DESC LIMIT 50")
        rows = cursor.fetchall()
        
        if not rows:
            ctk.CTkLabel(scroll_frame, text="No history found yet.", text_color="gray").pack(pady=20)
            return
            
        for row in rows:
            date_str, q_str, auth_str = row
            
            item_frame = ctk.CTkFrame(scroll_frame)
            item_frame.pack(fill="x", pady=5, padx=5)
            
            # Yazar varsa ekle, tarihi kısalt (sadece YYYY-MM-DD HH:MM)
            auth_disp = f" | Author: {auth_str}" if auth_str else ""
            disp_text = f"[{date_str[:16]}] {q_str}{auth_disp}"
            
            lbl = ctk.CTkLabel(item_frame, text=disp_text, anchor="w", justify="left")
            lbl.pack(side="left", padx=10, pady=10, fill="x", expand=True)
            
            # Load butonuna lambda ile o anki satırın verilerini bağlıyoruz
            load_btn = ctk.CTkButton(
                item_frame, 
                text="Load", 
                width=60, 
                command=lambda q=q_str, a=auth_str: self.load_history_item(q, a, history_win)
            )
            load_btn.pack(side="right", padx=10, pady=10)

    def load_history_item(self, query, author, window):
        """Seçilen geçmiş aramayı ana GUI'ye yükler."""
        # Arama çubuğunu temizle ve doldur
        self.search_entry.delete(0, "end")
        self.search_entry.insert(0, query)
        
        # Yazar çubuğunu temizle ve doldur
        self.author_entry.delete(0, "end")
        self.author_entry.insert(0, author if author else "")
        
        # Geçmiş penceresini kapat
        window.destroy()
        self.status_label.configure(text="Past search loaded. Ready to search.")
    # -----------------------------------------

    def update_slider_lbl(self, value):
        self.slider_val_label.configure(text=str(int(value)))

    def update_status(self, msg, progress_val):
        self.status_label.configure(text=msg)
        self.progress_bar.set(progress_val)

    def start_search_thread(self):
        query = self.search_entry.get().strip()
        if not query: return
        
        self.search_button.configure(state="disabled")
        self.word_btn.configure(state="disabled")
        self.ris_btn.configure(state="disabled")
        self.result_textbox.delete("0.0", "end")
        
        # Arama geçmişini kaydet
        self.conn.execute("INSERT INTO searches (date, query, author) VALUES (datetime('now', 'localtime'), ?, ?)", 
                          (query, self.author_entry.get().strip()))
        self.conn.commit()

        thread = threading.Thread(target=self.perform_search, 
                                  args=(query, self.author_entry.get().strip(), self.lang_dropdown.get(), int(self.article_slider.get())))
        thread.start()

    def perform_search(self, query, author, target_lang, num_articles):
        try:
            callback = lambda msg, val: self.after(0, self.update_status, msg, val)
            
            result = run_literature_review(query, author, target_lang, num_articles, callback)
            self.after(0, self.show_results, result["text"], result["articles"])
        except Exception as e:
            self.after(0, self.show_results, f"Error: {str(e)}", [])

    def show_results(self, text, articles):
        self.current_report = text
        self.current_articles = articles
        
        self.result_textbox.delete("0.0", "end")
        self.result_textbox.insert("0.0", text)
        self.search_button.configure(state="normal")
        self.status_label.configure(text="Process completed.")
        
        if articles:
            self.word_btn.configure(state="normal")
            self.ris_btn.configure(state="normal")

    def export_word(self):
        file_path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Document", "*.docx")], title="Save Report")
        if file_path:
            doc = Document()
            doc.add_heading('AI Literature Review Report', 0)
            doc.add_paragraph(self.current_report)
            doc.save(file_path)
            self.status_label.configure(text=f"Saved: {file_path}")

    def export_ris(self):
        file_path = filedialog.asksaveasfilename(defaultextension=".ris", filetypes=[("RIS File", "*.ris")], title="Save Citations for Zotero/Mendeley")
        if file_path:
            with open(file_path, "w", encoding="utf-8") as f:
                for a in self.current_articles:
                    f.write("TY  - JOUR\n")
                    f.write(f"TI  - {a['title']}\n")
                    for au in a['authors'].split(', '): f.write(f"AU  - {au}\n")
                    f.write(f"PY  - {a['year']}\n")
                    f.write(f"UR  - {a['doi']}\n")
                    f.write(f"C1  - {a['pmid']}\n")
                    f.write("ER  - \n\n")
            self.status_label.configure(text=f"RIS Saved: {file_path}")

if __name__ == "__main__":
    app = LiteratureApp()
    app.mainloop()